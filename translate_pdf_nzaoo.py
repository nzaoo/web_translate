import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import pdfplumber
from googletrans import Translator, LANGUAGES
from fpdf import FPDF
from docx import Document
import threading

# Chuẩn hóa tên ngôn ngữ
LANG_CODE_TO_NAME = {code: name.title() for code, name in LANGUAGES.items()}
LANG_NAME_TO_CODE = {name.title(): code for code, name in LANGUAGES.items()}

# Hàm dịch văn bản
translator = Translator()
def translate_text(text, source, target):
    # googletrans có thể giới hạn độ dài, nên chia nhỏ đoạn dài
    lines = text.split('\n')
    translated_lines = []
    for line in lines:
        if line.strip():
            try:
                translated = translator.translate(line, src=source, dest=target)
                translated_lines.append(translated.text)
            except Exception as e:
                translated_lines.append(f"[Lỗi dịch: {e}]")
        else:
            translated_lines.append("")
    return '\n'.join(translated_lines)

# Hàm trích xuất văn bản từ PDF
def extract_text_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        text = ""
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

# Hàm xuất ra PDF
def export_to_pdf(text, output_path):
    pdf_out = FPDF()
    pdf_out.add_page()
    pdf_out.set_font("Arial", size=12)
    for line in text.split('\n'):
        pdf_out.cell(0, 10, line, ln=True)
    pdf_out.output(output_path)

# Hàm xuất ra Word
def export_to_word(text, output_path):
    doc = Document()
    for line in text.split('\n'):
        doc.add_paragraph(line)
    doc.save(output_path)

# Hàm xử lý dịch và xuất file
def process_translate(input_path, source_lang, target_lang, output_path, output_type, update_status, show_texts):
    try:
        update_status("Đang trích xuất văn bản từ PDF...")
        text = extract_text_from_pdf(input_path)
        if not text.strip():
            update_status("Không tìm thấy văn bản trong file PDF!")
            show_texts("", "")
            return
        update_status(f"Đang dịch sang {LANG_CODE_TO_NAME.get(target_lang, target_lang)}...")
        translated = translate_text(text, source_lang, target_lang)
        update_status(f"Đang xuất ra file {output_type.upper()}...")
        if output_type == 'pdf':
            export_to_pdf(translated, output_path)
        else:
            export_to_word(translated, output_path)
        update_status(f"Hoàn thành! Đã lưu: {output_path}")
        show_texts(text, translated)
    except Exception as e:
        update_status("Lỗi: " + str(e))
        show_texts("", "")

# Giao diện Tkinter
class PDFTranslatorApp:
    def __init__(self, root):
        self.root = root
        root.title("Dịch PDF sang PDF/Word với Google Translate (Free)")
        root.geometry("1100x600")

        tk.Label(root, text="Chọn file PDF:").pack(anchor='w', padx=10, pady=(10,0))
        self.entry_input = tk.Entry(root, width=60)
        self.entry_input.pack(padx=10)
        tk.Button(root, text="Chọn file", command=self.select_file).pack(padx=10, pady=2)

        frame_lang = tk.Frame(root)
        frame_lang.pack(padx=10, pady=5, fill='x')
        tk.Label(frame_lang, text="Ngôn ngữ gốc:").grid(row=0, column=0, sticky='w')
        self.combo_source = ttk.Combobox(frame_lang, values=sorted(LANG_NAME_TO_CODE.keys()), width=25)
        self.combo_source.set('English')
        self.combo_source.grid(row=0, column=1, padx=5)
        tk.Label(frame_lang, text="Dịch sang:").grid(row=0, column=2, sticky='w', padx=(20,0))
        self.combo_target = ttk.Combobox(frame_lang, values=sorted(LANG_NAME_TO_CODE.keys()), width=25)
        self.combo_target.set('Vietnamese')
        self.combo_target.grid(row=0, column=3, padx=5)

        frame_type = tk.Frame(root)
        frame_type.pack(padx=10, pady=5, fill='x')
        tk.Label(frame_type, text="Định dạng xuất ra:").grid(row=0, column=0, sticky='w')
        self.output_type = tk.StringVar(value='pdf')
        tk.Radiobutton(frame_type, text="PDF", variable=self.output_type, value='pdf').grid(row=0, column=1)
        tk.Radiobutton(frame_type, text="Word (.docx)", variable=self.output_type, value='docx').grid(row=0, column=2)

        tk.Label(root, text="Tên file đầu ra:").pack(anchor='w', padx=10, pady=(10,0))
        self.entry_output = tk.Entry(root, width=60)
        self.entry_output.pack(padx=10)

        tk.Button(root, text="Dịch và Lưu", command=self.start_translate).pack(pady=10)
        self.label_status = tk.Label(root, text="")
        self.label_status.pack()

        # Split view
        split_frame = tk.Frame(root)
        split_frame.pack(fill='both', expand=True, padx=10, pady=10)
        tk.Label(split_frame, text="Nội dung gốc").pack(side='left', anchor='n')
        tk.Label(split_frame, text="Nội dung đã dịch").pack(side='right', anchor='n')
        self.text_original = tk.Text(split_frame, wrap='word', width=60)
        self.text_original.pack(side='left', fill='both', expand=True, padx=(0,5))
        self.text_translated = tk.Text(split_frame, wrap='word', width=60)
        self.text_translated.pack(side='right', fill='both', expand=True, padx=(5,0))

    def select_file(self):
        file_path = filedialog.askopenfilename(filetypes=[("PDF files", "*.pdf")])
        if file_path:
            self.entry_input.delete(0, tk.END)
            self.entry_input.insert(0, file_path)
            # Gợi ý tên file đầu ra
            base = os.path.splitext(os.path.basename(file_path))[0]
            ext = '.pdf' if self.output_type.get() == 'pdf' else '.docx'
            self.entry_output.delete(0, tk.END)
            self.entry_output.insert(0, base + '_translated' + ext)

    def start_translate(self):
        input_path = self.entry_input.get()
        if not input_path:
            messagebox.showwarning("Thiếu file", "Vui lòng chọn file PDF đầu vào.")
            return
        source_lang = LANG_NAME_TO_CODE.get(self.combo_source.get(), 'auto')
        target_lang = LANG_NAME_TO_CODE.get(self.combo_target.get(), 'vi')
        output_path = self.entry_output.get()
        if not output_path:
            messagebox.showwarning("Thiếu tên file", "Vui lòng nhập tên file đầu ra.")
            return
        output_type = self.output_type.get()
        threading.Thread(target=process_translate, args=(input_path, source_lang, target_lang, output_path, output_type, self.set_status, self.show_texts)).start()

    def set_status(self, msg):
        self.label_status.config(text=msg)

    def show_texts(self, original, translated):
        self.text_original.delete(1.0, tk.END)
        self.text_translated.delete(1.0, tk.END)
        self.text_original.insert(tk.END, original)
        self.text_translated.insert(tk.END, translated)

if __name__ == "__main__":
    root = tk.Tk()
    app = PDFTranslatorApp(root)
    root.mainloop() 